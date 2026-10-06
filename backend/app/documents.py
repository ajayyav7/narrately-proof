from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable

import fitz
from docx import Document
from openpyxl import load_workbook


@dataclass
class Passage:
    text: str
    location: str


def extract_document(filename: str, content: bytes) -> list[Passage]:
    """Extract readable text while preserving the best available source location."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise ValueError("Password-protected PDFs are not supported yet.")
            return [Passage(page.get_text("text").strip(), f"Page {number}")
                    for number, page in enumerate(pdf, start=1) if page.get_text("text").strip()]
    if suffix == ".docx":
        doc = Document(BytesIO(content))
        passages = [Passage(p.text.strip(), f"Paragraph {i}")
                    for i, p in enumerate(doc.paragraphs, start=1) if p.text.strip()]
        for table_no, table in enumerate(doc.tables, start=1):
            for row_no, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    passages.append(Passage(text, f"Table {table_no}, row {row_no}"))
        return passages
    if suffix == ".xlsx":
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        passages = []
        for sheet in book.worksheets:
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    passages.append(Passage(" | ".join(values), f"Sheet {sheet.title}, row {row_no}"))
        return passages
    raise ValueError("Unsupported format. Use PDF, DOCX, or XLSX.")


def split_claims(passages: Iterable[Passage]) -> list[Passage]:
    claims = []
    for passage in passages:
        for sentence in passage.text.replace("\n", " ").split(". "):
            text = sentence.strip().strip("•- ")
            # Keep meaningful report assertions, skip short headings and fragments.
            if len(text.split()) >= 6 and any(ch.isalpha() for ch in text):
                claims.append(Passage(text.rstrip(". "), passage.location))
    return claims[:500]
